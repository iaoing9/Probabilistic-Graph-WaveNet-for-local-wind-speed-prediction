import sys
import math
import traceback
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def choose_excel():
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title="Select QRF prediction-result Excel file",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
    )
    root.destroy()
    return path


def read_qrf_excel(file_path):
    if load_workbook is None:
        raise RuntimeError(
            "openpyxl is not installed.\n\n"
            "Please open Command Prompt and run:\n"
            "pip install openpyxl"
        )

    wb = load_workbook(file_path, data_only=True, read_only=True)
    try:
        ws = wb.active
        header_row = next(
            ws.iter_rows(min_row=1, max_row=1, values_only=True), None
        )
        if not header_row:
            raise ValueError("The Excel file does not contain a header row.")

        headers = [str(value).strip() if value is not None else "" for value in header_row]
        header_map = {name.casefold(): index for index, name in enumerate(headers) if name}
        rows = []
        for excel_row, row in enumerate(
            ws.iter_rows(min_row=2, max_col=ws.max_column, values_only=True), start=2
        ):
            if all(value is None for value in row):
                continue
            rows.append((excel_row, row))
    finally:
        wb.close()

    if not rows:
        raise ValueError("No data rows were found in the Excel file.")

    def find_column(candidates, label):
        for candidate in candidates:
            index = header_map.get(candidate.casefold())
            if index is not None:
                return index
        available = ", ".join(name for name in headers if name)
        raise ValueError(
            f"Required column for {label} was not found. Expected one of {candidates}. "
            f"Available columns: {available}"
        )

    time_index = find_column(("Target_time", "Time"), "time")
    observed_index = find_column(("Observed",), "observed values")
    mean_index = find_column(("Predicted_mean",), "predicted means")
    lower_index = find_column(("Lower_95PI", "Lower_PI"), "lower prediction interval")
    upper_index = find_column(("Upper_95PI", "Upper_PI"), "upper prediction interval")
    crps_index = find_column(("Sample_CRPS",), "sample-wise CRPS")
    horizon_index = header_map.get("horizon_hours_ahead")

    if horizon_index is not None:
        horizons = {
            str(row[horizon_index]).strip()
            for _, row in rows
            if horizon_index < len(row) and row[horizon_index] is not None
        }
        if len(horizons) > 1:
            raise ValueError(
                "This file contains multiple forecast horizons. Evaluate a single-horizon "
                "QRF prediction file so that metrics are not mixed across horizons."
            )

    time = []
    observed = []
    predicted_mean = []
    lower = []
    upper = []
    sample_crps = []

    required_indices = (
        observed_index,
        mean_index,
        lower_index,
        upper_index,
        crps_index,
    )

    for excel_row, row in rows:
        if any(index >= len(row) for index in required_indices):
            raise ValueError(f"Required prediction data are missing at Excel row {excel_row}.")

        try:
            y_value = float(row[observed_index])
            mean_value = float(row[mean_index])
            lower_value = float(row[lower_index])
            upper_value = float(row[upper_index])
            crps_value = float(row[crps_index])
        except (TypeError, ValueError):
            raise ValueError(
                f"Invalid or non-numeric data found in required columns at Excel row {excel_row}."
            ) from None

        if not all(
            math.isfinite(value)
            for value in (y_value, mean_value, lower_value, upper_value, crps_value)
        ):
            raise ValueError(
                f"NaN or Inf found in required columns at Excel row {excel_row}."
            )
        if lower_value > upper_value:
            raise ValueError(
                f"The lower prediction bound exceeds the upper bound at Excel row {excel_row}."
            )
        if crps_value < -1e-10:
            raise ValueError(
                f"CRPS must be non-negative at Excel row {excel_row}."
            )

        time.append(row[time_index])
        observed.append(y_value)
        predicted_mean.append(mean_value)
        lower.append(lower_value)
        upper.append(upper_value)
        sample_crps.append(max(crps_value, 0.0))

    return time, observed, predicted_mean, lower, upper, sample_crps


def calc_metrics(y, mu, lower, upper, sample_crps):
    lengths = {len(y), len(mu), len(lower), len(upper), len(sample_crps)}
    if len(lengths) != 1:
        raise ValueError("Observed values, predictions, intervals, and CRPS must have equal lengths.")

    n = len(y)
    if n == 0:
        raise ValueError("No valid samples are available.")

    errors = [yi - mi for yi, mi in zip(y, mu)]
    rmse = math.sqrt(math.fsum(error * error for error in errors) / n)
    mae = math.fsum(abs(error) for error in errors) / n

    smape_terms = []
    for yi, mi, error in zip(y, mu, errors):
        denominator = abs(yi) + abs(mi)
        smape_terms.append(0.0 if denominator == 0 else 2.0 * abs(error) / denominator)
    smape = math.fsum(smape_terms) / n

    y_mean = math.fsum(y) / n
    sse = math.fsum(error * error for error in errors)
    sst = math.fsum((yi - y_mean) ** 2 for yi in y)
    r2 = float("nan") if sst <= 0 else 1.0 - sse / sst

    covered = sum(1 for yi, lo, hi in zip(y, lower, upper) if lo <= yi <= hi)
    picp = covered / n

    y_range = max(y) - min(y)
    pinaw = (
        float("nan")
        if y_range <= 0
        else math.fsum(hi - lo for lo, hi in zip(lower, upper)) / n / y_range
    )

    crps = math.fsum(sample_crps) / n
    return [rmse, mae, smape, r2, picp, pinaw, crps]


def show_result(file_path, time, values):
    numeric_text = [
        "nan" if not math.isfinite(value) else f"{value:.4f}"
        for value in values
    ]
    result = "\t".join(numeric_text + ["/"])

    root = tk.Tk()
    root.title("Quantile-based metric evaluation")
    root.geometry("1120x270")
    root.resizable(True, False)

    tk.Label(
        root,
        text=(
            f"{Path(file_path).name}    |    {len(time)} rows    |    "
            f"Time: {time[0]} to {time[-1]}"
        ),
        font=("Segoe UI", 10),
    ).pack(pady=(15, 6))

    tk.Label(
        root,
        text="RMSE → MAE → sMAPE → R² → PICP → PINAW → CRPS → NLL",
        font=("Segoe UI", 10, "bold"),
    ).pack(pady=(0, 7))

    entry = tk.Entry(root, font=("Consolas", 12))
    entry.insert(0, result)
    entry.pack(fill="x", padx=18, pady=4)
    entry.select_range(0, tk.END)
    entry.focus_set()

    root.clipboard_clear()
    root.clipboard_append(result)
    root.update()

    tk.Label(
        root,
        text="Required columns verified by header names. NLL is reported as / for QRF.",
        font=("Segoe UI", 9),
    ).pack(pady=(7, 2))

    info = tk.Label(
        root,
        text="Result copied to clipboard. Paste directly into Excel with Ctrl+V.",
        font=("Segoe UI", 9),
    )
    info.pack(pady=(2, 8))

    def copy_again():
        root.clipboard_clear()
        root.clipboard_append(result)
        root.update()
        info.config(text="Result copied again.")

    frame = tk.Frame(root)
    frame.pack()
    tk.Button(frame, text="Copy result", width=14, command=copy_again).pack(side="left", padx=6)
    tk.Button(frame, text="Close", width=14, command=root.destroy).pack(side="left", padx=6)
    root.mainloop()


def main():
    try:
        file_path = sys.argv[1].strip('"') if len(sys.argv) >= 2 else choose_excel()
        if not file_path:
            return

        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File does not exist:\n{path}")

        time, observed, predicted_mean, lower, upper, sample_crps = read_qrf_excel(str(path))
        values = calc_metrics(observed, predicted_mean, lower, upper, sample_crps)
        show_result(str(path), time, values)

    except Exception as exc:
        detail = traceback.format_exc()
        try:
            root = tk.Tk()
            root.withdraw()
            messagebox.showerror(
                "Calculation failed",
                f"{exc}\n\nDetailed error:\n{detail}",
            )
            root.destroy()
        except Exception:
            print(detail)
            input("\nPress Enter to close...")


if __name__ == "__main__":
    main()
