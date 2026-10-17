import sys
import math
import traceback
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

Z95 = 1.96
EPS = 1e-12


def choose_excel():
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(
        title="Select prediction-result Excel file",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
    )
    root.destroy()
    return path


def read_excel_a_to_d(file_path):
    if load_workbook is None:
        raise RuntimeError(
            "openpyxl is not installed.\n\n"
            "Please open Command Prompt and run:\n"
            "pip install openpyxl"
        )

    wb = load_workbook(file_path, data_only=True, read_only=True)
    ws = wb.active

    headers = [ws.cell(1, c).value for c in range(1, 5)]

    rows = []
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=4, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(row)

    wb.close()

    if not rows:
        raise ValueError("No data rows were found in the Excel file.")

    t = []
    y = []
    mu = []
    var = []

    for idx, row in enumerate(rows, start=2):
        a, b, c, d = row
        try:
            b = float(b)
            c = float(c)
            d = float(d)
        except Exception:
            raise ValueError(
                f"Invalid/non-numeric data found in B-D at Excel row {idx}."
            )

        if not (math.isfinite(b) and math.isfinite(c) and math.isfinite(d)):
            raise ValueError(
                f"NaN/Inf found in B-D at Excel row {idx}."
            )
        if d <= 0:
            raise ValueError(
                f"Predicted variance must be positive at Excel row {idx}."
            )

        t.append(a)
        y.append(b)
        mu.append(c)
        var.append(max(d, EPS))

    return headers, t, y, mu, var


def calc_metrics(y, mu, var):
    n = len(y)
    err = [a - b for a, b in zip(y, mu)]

    rmse = math.sqrt(sum(e * e for e in err) / n)
    mae = sum(abs(e) for e in err) / n

    smape_terms = []
    for yi, mi, ei in zip(y, mu, err):
        den = abs(yi) + abs(mi)
        smape_terms.append(0.0 if den == 0 else 2.0 * abs(ei) / den)
    smape = sum(smape_terms) / n

    ybar = sum(y) / n
    sse = sum(e * e for e in err)
    sst = sum((yi - ybar) ** 2 for yi in y)
    r2 = float("nan") if sst <= 0 else 1.0 - sse / sst

    sigmas = [math.sqrt(v) for v in var]
    lowers = [m - Z95 * s for m, s in zip(mu, sigmas)]
    uppers = [m + Z95 * s for m, s in zip(mu, sigmas)]

    covered = sum(1 for yi, lo, hi in zip(y, lowers, uppers) if lo <= yi <= hi)
    picp = covered / n

    y_range = max(y) - min(y)
    pinaw = float("nan") if y_range <= 0 else (
        sum(hi - lo for lo, hi in zip(lowers, uppers)) / n / y_range
    )

    inv_sqrt_2pi = 1.0 / math.sqrt(2.0 * math.pi)
    inv_sqrt_pi = 1.0 / math.sqrt(math.pi)

    crps_sum = 0.0
    nll_sum = 0.0

    for yi, mi, vi, si in zip(y, mu, var, sigmas):
        z = (yi - mi) / si
        phi = inv_sqrt_2pi * math.exp(-0.5 * z * z)
        Phi = 0.5 * (1.0 + math.erf(z / math.sqrt(2.0)))

        crps_sum += si * (
            z * (2.0 * Phi - 1.0)
            + 2.0 * phi
            - inv_sqrt_pi
        )

        nll_sum += (
            0.5 * math.log(2.0 * math.pi * vi)
            + ((yi - mi) ** 2) / (2.0 * vi)
        )

    crps = crps_sum / n
    nll = nll_sum / n

    return [rmse, mae, smape, r2, picp, pinaw, crps, nll]


def show_result(file_path, t, values):
    result = "\t".join(
        "nan" if not math.isfinite(v) else f"{v:.4f}"
        for v in values
    )

    root = tk.Tk()
    root.title("Wind prediction 8 metrics")
    root.geometry("1020x245")
    root.resizable(True, False)

    tk.Label(
        root,
        text=f"{Path(file_path).name}    |    {len(t)} rows    |    Target_time: {t[0]}–{t[-1]}",
        font=("Segoe UI", 10)
    ).pack(pady=(16, 8))

    tk.Label(
        root,
        text="RMSE → MAE → sMAPE → R² → PICP → PINAW → CRPS → NLL",
        font=("Segoe UI", 10, "bold")
    ).pack(pady=(0, 8))

    e = tk.Entry(root, font=("Consolas", 12))
    e.insert(0, result)
    e.pack(fill="x", padx=18, pady=4)
    e.select_range(0, tk.END)
    e.focus_set()

    root.clipboard_clear()
    root.clipboard_append(result)
    root.update()

    info = tk.Label(
        root,
        text="Result copied to clipboard. Paste directly into Excel with Ctrl+V.",
        font=("Segoe UI", 9)
    )
    info.pack(pady=8)

    def copy_again():
        root.clipboard_clear()
        root.clipboard_append(result)
        root.update()
        info.config(text="Copied again.")

    frame = tk.Frame(root)
    frame.pack()

    tk.Button(frame, text="Copy result", width=14, command=copy_again).pack(side="left", padx=6)
    tk.Button(frame, text="Close", width=14, command=root.destroy).pack(side="left", padx=6)

    root.mainloop()


def main():
    try:
        if len(sys.argv) >= 2:
            file_path = sys.argv[1].strip('"')
        else:
            file_path = choose_excel()

        if not file_path:
            return

        p = Path(file_path)
        if not p.exists():
            raise FileNotFoundError(f"File does not exist:\n{p}")

        _, t, y, mu, var = read_excel_a_to_d(str(p))
        values = calc_metrics(y, mu, var)
        show_result(str(p), t, values)

    except Exception as exc:
        detail = traceback.format_exc()
        try:
            root = tk.Tk()
            root.withdraw()
            messagebox.showerror(
                "Calculation failed",
                f"{exc}\n\nDetailed error:\n{detail}"
            )
            root.destroy()
        except Exception:
            print(detail)
            input("\nPress Enter to close...")


if __name__ == "__main__":
    main()
